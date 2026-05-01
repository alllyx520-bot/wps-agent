# -*- coding: utf-8 -*-
"""Offline Word document builder — generates professional .docx files without WPS.
Requires: pip install python-docx
"""
from typing import Optional, Dict, List
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm, Emu, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
} if HAS_DOCX else {}


def _hex_to_rgb(hex_color: str):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def build_docx(structure: Dict, output_path: str) -> Dict:
    """Build a .docx from a JSON structure.
    structure = {
        "page": {"size": "A4", "orientation": "portrait",
                  "top_margin_cm": 2.54, "bottom_margin_cm": 2.54,
                  "left_margin_cm": 3.17, "right_margin_cm": 3.17},
        "default_font": {"name": "宋体", "size_pt": 12},
        "content": [
            {"type": "heading1", "text": "第一章 绪论"},
            {"type": "heading2", "text": "1.1 研究背景"},
            {"type": "body", "text": "正文内容...", "first_line_indent": True, "line_spacing": 1.5},
            {"type": "page_break"},
            {"type": "table", "headers": ["列1","列2"], "rows": [["A","B"]]},
        ]
    }
    """
    if not HAS_DOCX:
        return {"error": "python-docx not installed. Run: pip install python-docx"}
    doc = Document()

    # Page setup
    page_cfg = structure.get("page", {})
    section = doc.sections[0]
    section.page_width = Cm(21.0)   # A4
    section.page_height = Cm(29.7)
    if page_cfg:
        if page_cfg.get("orientation") == "landscape":
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Cm(page_cfg.get("top_margin_cm", 2.54))
        section.bottom_margin = Cm(page_cfg.get("bottom_margin_cm", 2.54))
        section.left_margin = Cm(page_cfg.get("left_margin_cm", 3.17))
        section.right_margin = Cm(page_cfg.get("right_margin_cm", 3.17))

    # Default font
    df = structure.get("default_font", {})
    style = doc.styles["Normal"]
    style.font.name = df.get("name", "宋体")
    style.font.size = Pt(df.get("size_pt", 12))
    style.element.rPr.rFonts.set(qn("w:eastAsia"), df.get("name", "宋体"))
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    stats = {"paragraphs": 0, "tables": 0, "page_breaks": 0}
    for item in structure.get("content", []):
        item_type = item.get("type", "body")
        text = item.get("text", "")

        if item_type == "page_break":
            doc.add_page_break()
            stats["page_breaks"] += 1
            continue

        if item_type.startswith("table"):
            headers = item.get("headers", [])
            rows_data = item.get("rows", [])
            if not rows_data and not headers:
                continue
            tbl = doc.add_table(rows=1 + len(rows_data), cols=len(headers) or (len(rows_data[0]) if rows_data else 1))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = "Table Grid"
            if headers:
                for j, h in enumerate(headers):
                    cell = tbl.rows[0].cells[j]
                    cell.text = str(h)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
            for i, row_data in enumerate(rows_data):
                for j, val in enumerate(row_data):
                    tbl.rows[i + 1].cells[j].text = str(val)
            stats["tables"] += 1
            continue

        p = doc.add_paragraph()
        if item_type in ("heading1", "heading2", "heading3"):
            level = int(item_type[-1])
            p.style = doc.styles[f"Heading {level}"]
            p.add_run(text)
        elif item_type == "body":
            run = p.add_run(text)
            if item.get("font"):
                f = item["font"]
                run.font.name = f.get("name", df.get("name", "宋体"))
                run.font.size = Pt(f.get("size_pt", df.get("size_pt", 12)))
                if f.get("bold"):
                    run.bold = True
                run.element.rPr.rFonts.set(qn("w:eastAsia"), f.get("name", df.get("name", "宋体")))
            if item.get("first_line_indent") or item.get("first_line_indent_chars"):
                chars = item.get("first_line_indent_chars", 2)
                p.paragraph_format.first_line_indent = Pt(chars * (item.get("font", {}).get("size_pt", df.get("size_pt", 12))))
            if item.get("line_spacing"):
                p.paragraph_format.line_spacing = item["line_spacing"]
            if item.get("alignment"):
                p.alignment = ALIGN_MAP.get(item["alignment"], WD_ALIGN_PARAGRAPH.LEFT)
        else:
            run = p.add_run(text)
        stats["paragraphs"] += 1

    doc.save(output_path)
    return {"output": output_path, "stats": stats}


def build_cover_page(lines: List[Dict], output_path: str) -> Dict:
    """Generate a standalone cover page docx.
    lines = [{"text": "Title", "font_name": "黑体", "font_size_pt": 26, "bold": True, "alignment": "center", "space_before_pt": 120}, ...]
    """
    if not HAS_DOCX:
        return {"error": "python-docx not installed"}
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for ln in lines:
        text = ln.get("text", "")
        if not text.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(ln.get("space_before_pt", 6))
            continue
        p = doc.add_paragraph()
        run = p.add_run(text)
        font_name = ln.get("font_name", "宋体")
        run.font.name = font_name
        run.font.size = Pt(ln.get("font_size_pt", 14))
        run.bold = ln.get("bold", False)
        run.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        p.alignment = ALIGN_MAP.get(ln.get("alignment", "center"), WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.space_before = Pt(ln.get("space_before_pt", 0))
        p.paragraph_format.space_after = Pt(ln.get("space_after_pt", 6))
    doc.save(output_path)
    return {"output": output_path, "lines": len(lines)}
